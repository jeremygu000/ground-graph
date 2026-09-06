"""Document parsers for supported formats.

Each parser returns a structured ``ParsedContent`` that the chunker consumes.
Parsers may raise ``UnsupportedFormatError`` to signal a typed reason code
when the format cannot be processed.
"""

from __future__ import annotations

import io
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from enum import StrEnum
from typing import ClassVar

import docx  # pyright: ignore[reportMissingTypeStubs]
from ebooklib import epub  # pyright: ignore[reportMissingTypeStubs]
from pypdf import PdfReader  # pyright: ignore[reportMissingTypeStubs]


def _metadata_default() -> dict[str, object]:
    return {}


def _headings_default() -> list[tuple[int, str]]:
    return []


def _protected_ranges_default() -> list[tuple[int, int]]:
    return []


@dataclass
class ParsedContent:
    title: str
    body: str
    media_type: str
    metadata: dict[str, object] = field(default_factory=_metadata_default)
    headings: list[tuple[int, str]] = field(default_factory=_headings_default)
    code_blocks: list[tuple[int, int]] = field(default_factory=_protected_ranges_default)
    tables: list[tuple[int, int]] = field(default_factory=_protected_ranges_default)
    list_items: list[tuple[int, int]] = field(default_factory=_protected_ranges_default)


class BaseParser(ABC):
    @property
    @abstractmethod
    def media_type(self) -> str: ...

    @abstractmethod
    def parse(self, content: bytes) -> ParsedContent: ...


class TextParser(BaseParser):
    @property
    def media_type(self) -> str:
        return "text/plain"

    def parse(self, content: bytes) -> ParsedContent:
        text = content.decode("utf-8", errors="replace")
        first_line = text.split("\n", maxsplit=1)[0].strip()
        title = first_line[:120] if first_line else "untitled"
        return ParsedContent(
            title=title,
            body=text,
            media_type=self.media_type,
            metadata={"charset": "utf-8"},
        )


class MarkdownParser(BaseParser):
    @property
    def media_type(self) -> str:
        return "text/markdown"

    def parse(self, content: bytes) -> ParsedContent:
        text = content.decode("utf-8", errors="replace")
        title, body = self._strip_frontmatter(text)
        headings = self._extract_headings(body)
        code_blocks = self._extract_code_blocks(body)
        tables = self._extract_tables(body)
        list_items = self._extract_list_items(body)
        return ParsedContent(
            title=title or "untitled",
            body=body,
            media_type=self.media_type,
            metadata={},
            headings=headings,
            code_blocks=code_blocks,
            tables=tables,
            list_items=list_items,
        )

    def _strip_frontmatter(self, text: str) -> tuple[str, str]:
        fm_match = re.match(r"^---\s*\n.*?\n---\s*\n", text, re.DOTALL)
        if not fm_match:
            title = self._title_from_first_header(text)
            return title, text
        rest = text[fm_match.end() :]
        title = self._title_from_first_header(rest)
        return title, text

    def _title_from_first_header(self, text: str) -> str:
        m = re.match(r"^#\s+(.+)$", text, re.MULTILINE)
        return m.group(1).strip() if m else ""

    def _extract_headings(self, body: str) -> list[tuple[int, str]]:
        results: list[tuple[int, str]] = []
        for m in re.finditer(r"^(#{1,6})\s+(.+)$", body, re.MULTILINE):
            level = len(m.group(1))
            results.append((level, m.group(2).strip()))
        return results

    def _extract_code_blocks(self, body: str) -> list[tuple[int, int]]:
        results: list[tuple[int, int]] = []
        pattern = re.compile(r"^```", re.MULTILINE)
        starts: list[int] = []
        for m in pattern.finditer(body):
            line_no = body.count("\n", 0, m.start()) + 1
            starts.append(line_no)
        for i in range(0, len(starts) - 1, 2):
            results.append((starts[i], starts[i + 1]))
        return results

    def _extract_tables(self, body: str) -> list[tuple[int, int]]:
        results: list[tuple[int, int]] = []
        lines = body.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if "|" in line and not line.startswith("```"):
                sep_line = i + 1
                if sep_line < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[sep_line].strip()):
                    end = sep_line + 1
                    if end < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[end].strip()):
                        end += 1
                    while end <= len(lines) and lines[end - 1].strip().startswith("|"):
                        end += 1
                    results.append((i + 1, end - 1))
                    i = end
                    continue
            i += 1
        return results

    def _extract_list_items(self, body: str) -> list[tuple[int, int]]:
        results: list[tuple[int, int]] = []
        lines = body.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if re.match(r"^[-*+]\s+\S|^\d+\.\s+\S", line):
                start = i + 1
                indent = len(line) - len(line.lstrip())
                j = i + 1
                while j < len(lines):
                    next_line = lines[j].strip()
                    if not next_line:
                        j += 1
                        continue
                    next_indent = len(lines[j]) - len(lines[j].lstrip())
                    if next_indent > indent and (
                        next_line.startswith("-")
                        or next_line.startswith("*")
                        or next_line.startswith("+")
                        or re.match(r"^\d+\.", next_line)
                    ):
                        break
                    if next_indent <= indent and next_line and not next_line.startswith(" "):
                        break
                    j += 1
                results.append((start, j))
                i = j
                continue
            i += 1
        return results


class HtmlParser(BaseParser):
    @property
    def media_type(self) -> str:
        return "text/html"

    def parse(self, content: bytes) -> ParsedContent:
        text = content.decode("utf-8", errors="replace")
        title = self._extract_title(text)
        body = self._normalize_to_markdown(text)
        headings = self._extract_headings(body)
        code_blocks = self._extract_code_blocks(body)
        tables = self._extract_tables(body)
        list_items = self._extract_list_items(body)
        return ParsedContent(
            title=title or "untitled",
            body=body,
            media_type=self.media_type,
            metadata={},
            headings=headings,
            code_blocks=code_blocks,
            tables=tables,
            list_items=list_items,
        )

    def _extract_title(self, html: str) -> str:
        m = re.search(r"<title[^>]*>(.+?)</title>", html, re.IGNORECASE | re.DOTALL)
        if m:
            return m.group(1).strip()
        m = re.search(r"<h1[^>]*>(.+?)</h1>", html, re.IGNORECASE | re.DOTALL)
        return m.group(1).strip() if m else ""

    STRIP_TAGS_RE = re.compile(r"<[^>]+>")
    WHITESPACE_RE = re.compile(r"[ \t]+")

    def _normalize_to_markdown(self, html: str) -> str:
        def _inner_text(tag: str) -> str:
            return re.sub(r"<[^>]+>", "", tag).strip()

        def _replace_heading(m: re.Match[str]) -> str:
            level = int(m.group(1))
            inner = m.group(2)
            text = re.sub(r"<[^>]+>", "", inner).strip()
            return "\n" + "#" * level + " " + text + "\n"

        html = re.sub(
            r"<h([1-6])(?:\s[^>]*)?>(.*?)</h\1>",
            _replace_heading,
            html,
            flags=re.DOTALL | re.IGNORECASE,
        )
        html = re.sub(
            r"<pre><code([^>]*)>(.*?)</code></pre>",
            lambda m: f"\n```\n{_inner_text(m.group(2))}\n```\n",
            html,
            flags=re.DOTALL | re.IGNORECASE,
        )
        html = re.sub(
            r"<code([^>]*)>(.*?)</code>",
            lambda m: f"`{_inner_text(m.group(2))}`",
            html,
            flags=re.DOTALL | re.IGNORECASE,
        )

        def _replace_table(m: re.Match[str]) -> str:
            inner = m.group(0)
            rows = re.findall(r"<tr[^>]*>(.*?)</tr>", inner, re.DOTALL | re.IGNORECASE)
            if not rows:
                return ""
            md_rows: list[str] = []
            for row_html in rows:
                cells = re.findall(
                    r"<t[hd][^>]*>(.*?)</t[hd]>", row_html, re.DOTALL | re.IGNORECASE
                )
                if cells:
                    text_cells = [re.sub(r"<[^>]+>", "", c).strip() for c in cells]
                    md_rows.append("| " + " | ".join(text_cells) + " |")
            if not md_rows:
                return ""
            col_count = md_rows[0].count("|") - 1
            sep = "| " + " | ".join(["---"] * col_count) + " |"
            if len(md_rows) > 1:
                return "\n" + "\n".join([md_rows[0], sep, *md_rows[1:]]) + "\n"
            return "\n" + md_rows[0] + "\n" + sep + "\n"

        html = re.sub(
            r"<table[^>]*>.*?</table>", _replace_table, html, flags=re.DOTALL | re.IGNORECASE
        )

        def _replace_list_item(m: re.Match[str]) -> str:
            inner = re.sub(r"<[^>]+>", "", m.group(0)).strip()
            return "\n- " + inner + "\n"

        html = re.sub(
            r"<li[^>]*>.*?</li>", _replace_list_item, html, flags=re.DOTALL | re.IGNORECASE
        )

        text = self.STRIP_TAGS_RE.sub(" ", html)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = self.WHITESPACE_RE.sub(" ", text)
        return text.strip()

    def _extract_headings(self, body: str) -> list[tuple[int, str]]:
        results: list[tuple[int, str]] = []
        for m in re.finditer(r"^(#{1,6})\s+(.+)$", body, re.MULTILINE):
            level = len(m.group(1))
            results.append((level, m.group(2).strip()))
        return results

    def _extract_code_blocks(self, body: str) -> list[tuple[int, int]]:
        results: list[tuple[int, int]] = []
        pattern = re.compile(r"^```", re.MULTILINE)
        starts: list[int] = []
        for m in pattern.finditer(body):
            line_no = body.count("\n", 0, m.start()) + 1
            starts.append(line_no)
        for i in range(0, len(starts) - 1, 2):
            results.append((starts[i], starts[i + 1]))
        return results

    def _extract_tables(self, body: str) -> list[tuple[int, int]]:
        results: list[tuple[int, int]] = []
        lines = body.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if "|" in line and not line.startswith("```"):
                sep_line = i + 1
                if sep_line < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[sep_line].strip()):
                    end = sep_line + 1
                    if end < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[end].strip()):
                        end += 1
                    while end <= len(lines) and lines[end - 1].strip().startswith("|"):
                        end += 1
                    results.append((i + 1, end - 1))
                    i = end
                    continue
            i += 1
        return results

    def _extract_list_items(self, body: str) -> list[tuple[int, int]]:
        results: list[tuple[int, int]] = []
        lines = body.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if re.match(r"^[-*+]\s+\S|^\d+\.\s+\S", line):
                start = i + 1
                indent = len(line) - len(line.lstrip())
                j = i + 1
                while j < len(lines):
                    next_line = lines[j].strip()
                    if not next_line:
                        j += 1
                        continue
                    next_indent = len(lines[j]) - len(lines[j].lstrip())
                    if next_indent > indent and (
                        next_line.startswith("-")
                        or next_line.startswith("*")
                        or next_line.startswith("+")
                        or re.match(r"^\d+\.", next_line)
                    ):
                        break
                    if next_indent <= indent and next_line and not next_line.startswith(" "):
                        break
                    j += 1
                results.append((start, j))
                i = j
                continue
            i += 1
        return results


class UnsupportedReason(StrEnum):
    UNSUPPORTED_MEDIA_TYPE = "unsupported_media_type"
    PDF_ENCRYPTED = "pdf_encrypted"
    PDF_SCANNED_ONLY = "pdf_scanned_only"
    PDF_NO_TEXT_LAYER = "pdf_no_text_layer"
    DOCX_PARSE_ERROR = "docx_parse_error"
    EPUB_PARSE_ERROR = "epub_parse_error"
    EMPTY_CONTENT = "empty_content"
    PARSE_ERROR = "parse_error"


@dataclass
class UnsupportedFormatError(Exception):
    reason: UnsupportedReason
    detail: str = ""

    def __str__(self) -> str:
        if self.detail:
            return f"{self.reason.value}: {self.detail}"
        return self.reason.value


class PdfParser(BaseParser):
    @property
    def media_type(self) -> str:
        return "application/pdf"

    def parse(self, content: bytes) -> ParsedContent:  # type: ignore[return-type]
        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception as exc:
            raise UnsupportedFormatError(UnsupportedReason.PARSE_ERROR, str(exc)) from exc

        if reader.is_encrypted:
            raise UnsupportedFormatError(UnsupportedReason.PDF_ENCRYPTED)

        lines: list[str] = []
        code_blocks: list[tuple[int, int]] = []
        in_code = False
        code_start = 0

        for page in reader.pages:
            text = page.extract_text()
            if not text or not text.strip():
                continue
            page_lines = text.splitlines()
            for line in page_lines:
                line_no = len(lines) + 1
                stripped = line.strip()
                if stripped.startswith("```"):
                    if not in_code:
                        in_code = True
                        code_start = line_no
                    else:
                        in_code = False
                        code_blocks.append((code_start, line_no))
                    continue
                if not in_code:
                    lines.append(line + "\n")

        body = "".join(lines)

        if not body.strip():
            raise UnsupportedFormatError(UnsupportedReason.PDF_NO_TEXT_LAYER)

        return ParsedContent(
            title=self._title_from_first_line(body) or "PDF page 1",
            body=body,
            media_type=self.media_type,
            metadata={"page_count": len(reader.pages)},
            code_blocks=code_blocks,
        )

    def _title_from_first_line(self, body: str) -> str:
        first = body.split("\n", maxsplit=1)[0].strip()
        return first[:120] if first else ""


class DocxParser(BaseParser):
    @property
    def media_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def parse(self, content: bytes) -> ParsedContent:
        try:
            document = docx.document.Document(io.BytesIO(content))
        except Exception as exc:
            raise UnsupportedFormatError(UnsupportedReason.DOCX_PARSE_ERROR, str(exc)) from exc

        lines: list[str] = []
        code_blocks: list[tuple[int, int]] = []
        in_code = False
        code_start = 0

        for para in document.paragraphs:
            text = para.text
            if not text.strip():
                lines.append("\n")
                continue
            line_no = len(lines) + 1
            stripped = text.strip()
            if stripped.startswith("```"):
                if not in_code:
                    in_code = True
                    code_start = line_no
                else:
                    in_code = False
                    code_blocks.append((code_start, line_no))
                lines.append(text + "\n")
            elif in_code:
                lines.append(text + "\n")
            elif para.style and "Code" in para.style.name:
                lines.append(f"```\n{text}\n```\n")
            else:
                lines.append(text + "\n")

        body = "".join(lines)

        row_count = sum(len(table.rows) for table in document.tables)
        tables: list[tuple[int, int]] = [(0, row_count)] if row_count > 0 else []

        title = document.core_properties.title or ""
        if not title:
            first = body.split("\n", maxsplit=1)[0].strip()
            title = first[:120] if first else "untitled"

        if not body.strip():
            raise UnsupportedFormatError(UnsupportedReason.EMPTY_CONTENT)

        return ParsedContent(
            title=title,
            body=body,
            media_type=self.media_type,
            metadata={
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            },
            code_blocks=code_blocks,
            tables=tables,
        )


class EpubParser(BaseParser):
    MAX_HEADING_LEVEL = 6

    @property
    def media_type(self) -> str:
        return "application/epub+zip"

    def parse(self, content: bytes) -> ParsedContent:
        try:
            book = epub.read_epub(io.BytesIO(content))
        except Exception as exc:
            raise UnsupportedFormatError(UnsupportedReason.EPUB_PARSE_ERROR, str(exc)) from exc

        title_meta = book.get_metadata("DC", "title")
        if title_meta:
            title_str = title_meta[0] if isinstance(title_meta, list) else str(title_meta)
        else:
            title_str = "untitled"

        body_parts, code_blocks, headings = self._extract_epub_content(book)

        body = "".join(body_parts)

        if not body.strip():
            raise UnsupportedFormatError(UnsupportedReason.EMPTY_CONTENT)

        return ParsedContent(
            title=title_str[:120] if title_str else "untitled",
            body=body,
            media_type=self.media_type,
            metadata={"chapter_count": len(list(book.get_items()))},
            headings=headings,
            code_blocks=code_blocks,
        )

    def _extract_epub_content(
        self, book: object
    ) -> tuple[list[str], list[tuple[int, int]], list[tuple[int, str]]]:
        body_parts: list[str] = []
        code_blocks: list[tuple[int, int]] = []
        headings: list[tuple[int, str]] = []
        in_code = False
        code_start = 0
        line_offset = 0

        for item in book.get_items():
            if item.get_type() != epub.ITEM_DOCUMENT:
                continue
            try:
                item_content = item.get_content().decode("utf-8", errors="replace")
            except Exception:
                continue

            text = self._strip_html_tags(item_content)
            if not text.strip():
                continue

            lines = text.splitlines()
            for line in lines:
                line_no = line_offset + 1
                stripped = line.strip()
                if stripped.startswith("#"):
                    level = len(stripped) - len(stripped.lstrip("#"))
                    heading_text = stripped.lstrip("#").strip()
                    if 1 <= level <= self.MAX_HEADING_LEVEL:
                        headings.append((level, heading_text))
                    body_parts.append(line + "\n")
                elif stripped.startswith("```"):
                    if not in_code:
                        in_code = True
                        code_start = line_no
                    else:
                        in_code = False
                        code_blocks.append((code_start, line_no))
                    body_parts.append(line + "\n")
                else:
                    body_parts.append(line + "\n")

                line_offset = line_no

        return body_parts, code_blocks, headings

    def _strip_html_tags(self, html: str) -> str:
        text = re.sub(r"<[^>]+>", "", html)
        text = re.sub(r"\s+", " ", text)
        return text.strip()


class ParserRegistry:
    _parsers: ClassVar[dict[str, BaseParser]] = {}

    @classmethod
    def register(cls, parser: BaseParser) -> None:
        cls._parsers[parser.media_type] = parser

    @classmethod
    def get(cls, media_type: str) -> BaseParser | None:
        return cls._parsers.get(media_type)

    @classmethod
    def get_with_reason(cls, media_type: str) -> BaseParser:
        parser = cls._parsers.get(media_type)
        if parser is None:
            raise UnsupportedFormatError(UnsupportedReason.UNSUPPORTED_MEDIA_TYPE, media_type)
        return parser


for _parser_cls in (
    TextParser,
    MarkdownParser,
    HtmlParser,
    PdfParser,
    DocxParser,
    EpubParser,
):
    ParserRegistry.register(_parser_cls())
